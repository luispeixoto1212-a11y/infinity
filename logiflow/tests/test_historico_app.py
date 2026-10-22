"""Testes da interface do histórico (versão 1.3), sem abrir o navegador."""

from __future__ import annotations

import datetime as dt
import io
import time
from pathlib import Path

import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from logiflow import conservacao as R
from logiflow import copias
from logiflow import historico as H
from logiflow import persistencia as P
from logiflow.exemplo import DATA_REFERENCIA_EXEMPLO, FICHEIRO_EXEMPLO_DIA2_XLSX, FICHEIRO_EXEMPLO_XLSX
from logiflow.exportacao import exportar_dados_limpos, exportar_relatorio

from .conftest import CABECALHO, csv_bytes, guardar, processar

RAIZ = Path(__file__).resolve().parent.parent
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DIA1, DIA2 = FICHEIRO_EXEMPLO_XLSX.read_bytes(), FICHEIRO_EXEMPLO_DIA2_XLSX.read_bytes()
SEPARADORES = ["Encomendas", "Exceções (18)", "Resumo", "Desempenho", "Qualidade dos dados (11)", "Exportar",
               "Histórico"]


# Relógio fixo ------------------------------------------------------------------------------------------
# A aplicação usa o dia de hoje (datetime.date.today) como data de referência predefinida, e os atrasos
# dependem dela. Nos testes que contam atrasos, o relógio fica no dia para que os dados do exemplo foram
# feitos; a aplicação não é alterada. Como date.today() equivale a date.fromtimestamp(time.time()), basta
# deslocar time.time: as horas continuam a passar e as datas continuam a ser do tipo normal.

_TIME_REAL = time.time
_desvio_s = 0.0  # diferente de 0 só durante um teste com o relógio fixo


def _time_deslocado() -> float:
    return _TIME_REAL() + _desvio_s


@pytest.fixture
def hoje_e_o_dia_do_exemplo(monkeypatch):
    """O relógio passa a marcar 20/10/2026 (ao meio-dia, e a avançar) só durante o teste, seja qual for a
    data do sistema. No fim, o desvio volta a 0: se algum módulo tiver guardado a função, dá a hora real."""
    global _desvio_s
    _desvio_s = dt.datetime.combine(DATA_REFERENCIA_EXEMPLO, dt.time(12)).timestamp() - _TIME_REAL()
    monkeypatch.setattr(time, "time", _time_deslocado)
    try:
        assert dt.date.today() == DATA_REFERENCIA_EXEMPLO
        yield DATA_REFERENCIA_EXEMPLO
    finally:
        _desvio_s = 0.0


def _app() -> AppTest:
    return AppTest.from_file(str(RAIZ / "app.py"), default_timeout=60).run()


def _carregar(at: AppTest, nome: str, conteudo: bytes, tipo: str = XLSX) -> AppTest:
    at.file_uploader(key="ficheiro").set_value((nome, conteudo, tipo)).run()
    return at


def _mensagens(at: AppTest) -> list[str]:
    return [c.value for c in at.caption if c.value.startswith(":material/history")]


def _lista(at: AppTest):
    return next(d.value for d in at.dataframe if "Importada em" in d.value.columns)


def _estado_dos_outros_separadores(at: AppTest):
    metricas = [(m.label, m.value) for m in at.metric]
    return metricas, [t.label for t in at.tabs], [c.value for c in at.caption if "encomendas com exceções" in c.value]


def _base() -> Path:
    return P.caminho_base_dados()


# Gravação automática e mensagens ----------------------------------------------------------------------

def test_ficheiro_importado_fica_no_historico(pasta_dados_temporaria, hoje_e_o_dia_do_exemplo):
    at = _carregar(_app(), "semana42.xlsx", DIA1)
    assert not at.exception and not at.error
    assert _mensagens(at) == [":material/history: Guardada no histórico como importação n\\.º 1 "
                              "\\(concluída com avisos\\)\\."]
    assert [t.label.split(" ", 1)[1] for t in at.tabs] == SEPARADORES
    lista = _lista(at)
    assert list(lista["Ficheiro"]) == ["semana42.xlsx"] and list(lista["Estado"]) == ["Concluída com avisos"]
    assert at.date_input(key="data_referencia").value == hoje_e_o_dia_do_exemplo
    assert (lista.loc[0, "Linhas"], lista.loc[0, "Entregues"], lista.loc[0, "Atrasadas"]) == ("60", "41", "8")
    # Interagir com a aplicação não volta a gravar.
    at.number_input(key="prazo_padrao").set_value(3).run()
    assert H.listar_importacoes(_base())[1] == 1


def test_mesmo_ficheiro_noutra_sessao_e_indicado_como_repetido(pasta_dados_temporaria):
    _carregar(_app(), "a.xlsx", DIA1)
    at = _carregar(_app(), "b.xlsx", DIA1)
    assert "já tinha sido importado" in _mensagens(at)[0] and "importação n\\.º 1" in _mensagens(at)[0]
    assert H.listar_importacoes(_base())[1] == 1


def test_voltar_a_carregar_o_mesmo_ficheiro_na_mesma_sessao(pasta_dados_temporaria):
    at = _carregar(_app(), "dia1.xlsx", DIA1)
    _carregar(at, "dia2.xlsx", DIA2)
    _carregar(at, "dia1.xlsx", DIA1)
    assert "já tinha sido importado" in _mensagens(at)[0] and "importação n\\.º 1" in _mensagens(at)[0]
    assert H.listar_importacoes(_base())[1] == 2


def test_opcao_desligada_nao_grava_nada(pasta_dados_temporaria, hoje_e_o_dia_do_exemplo):
    at = _app()
    at.toggle(key="guardar_historico").set_value(False).run()
    _carregar(at, "a.xlsx", DIA1)
    assert any("não foi guardada no histórico" in m for m in _mensagens(at))
    assert H.listar_importacoes(_base())[1] == 0
    assert at.date_input(key="data_referencia").value == hoje_e_o_dia_do_exemplo
    assert [(m.label, m.value) for m in at.metric][:5] == [("Encomendas", "60"), ("Entregues", "41"),
                                                          ("Pendentes", "11"), ("Atrasadas", "8"),
                                                          ("Linhas com erros", "9")]


def test_ficheiro_ilegivel_fica_registado_como_falhado(pasta_dados_temporaria):
    at = _carregar(_app(), "vazio.csv", b"", "text/csv")
    assert at.error and "Não foi possível importar" in at.error[0].value
    assert "importação falhada" in _mensagens(at)[0]
    lista, total = H.listar_importacoes(_base())
    assert total == 1 and lista.loc[0, "estado"] == H.FALHADA and lista.loc[0, "nome_ficheiro"] == "vazio.csv"


def test_exemplo_so_e_guardado_a_pedido(pasta_dados_temporaria):
    at = _app()
    at.button(key="usar_exemplo_inicio").click().run()
    assert "não é guardado automaticamente" in _mensagens(at)[0]
    assert H.listar_importacoes(_base())[1] == 0
    at.button(key="hist_guardar_exemplo").click().run()
    assert not at.exception and any("Guardada no histórico" in s.value for s in at.success)
    lista = H.listar_importacoes(_base())[0]
    assert list(lista["origem"]) == [H.ORIGEM_EXEMPLO] and lista.loc[0, "data_referencia"] == DATA_REFERENCIA_EXEMPLO.isoformat()


def test_falha_ao_gravar_e_comunicada_e_a_aplicacao_continua(pasta_dados_temporaria, monkeypatch):
    original = H._inserir_fotografia
    monkeypatch.setattr(H, "_inserir_fotografia", lambda *a: (_ for _ in ()).throw(RuntimeError("x")))
    at = _carregar(_app(), "a.xlsx", DIA1)
    assert not at.exception
    aviso = at.warning[0].value
    assert "Não foi possível guardar no histórico" in aviso and "continua a funcionar" in aviso
    assert [(m.label, m.value) for m in at.metric][0] == ("Encomendas", "60")
    monkeypatch.setattr(H, "_inserir_fotografia", original)
    at.button(key="hist_tentar_de_novo").click().run()
    assert "Guardada no histórico como importação n\\.º 2" in _mensagens(at)[0]


# Detalhes, comparação e ecrã inicial ------------------------------------------------------------------

def test_detalhes_e_comparacao_com_a_importacao_anterior(pasta_dados_temporaria):
    at = _carregar(_app(), "dia1.xlsx", DIA1)
    _carregar(at, "dia2.xlsx", DIA2)
    assert list(_lista(at)["Ficheiro"]) == ["dia2.xlsx", "dia1.xlsx"]
    assert at.selectbox(key="hist_detalhe").value == 2
    detalhe = next(m.value for m in at.markdown if m.value.startswith("**Estado:**"))
    assert "dia2\\.xlsx" in detalhe and "Importação anterior concluída:** n.º 1" in detalhe
    at.button(key="hist_comparar_anterior").click().run()
    assert not at.exception
    assert at.selectbox(key="hist_detalhe").value == 2  # os detalhes continuam na mesma importação
    resumo = next(m.value for m in at.markdown if m.value.startswith("Nova:"))
    assert resumo == ("Nova: **5** · Alterada: **7** · Sem alterações: **46** · "
                      "Não encontrada na importação atual: **3**")
    assert any("não quer dizer que a encomenda foi entregue nem cancelada" in c.value for c in at.caption)
    at.multiselect(key="hist_cmp_situacoes").set_value(["Nova"]).run()
    comparacao = next(d.value for d in at.dataframe if "Alterações" in d.value.columns)
    assert list(comparacao["Situação"]) == ["Nova"] * 5
    assert set(comparacao["Linha (anterior)"]) == {"—"}  # sem linha na anterior: «—», nunca «None»
    at.toggle(key="hist_mostrar_fotografia").set_value(True).run()
    guardadas = next(d.value for d in at.dataframe if "Erros" in d.value.columns and "Linha" in d.value.columns)
    assert len(guardadas) == 62


def test_ecra_inicial_da_acesso_ao_historico(pasta_dados_temporaria):
    guardar(_base_preparada(), csv_bytes(CABECALHO, "A1;Um;01/10/2026;Porto;Pendente"), "antiga.csv")
    at = _app()
    assert not at.exception and "Histórico de importações e cópias de segurança" in [e.label for e in at.expander]
    assert list(_lista(at)["Ficheiro"]) == ["antiga.csv"]


def _base_preparada() -> Path:
    caminho = _base()
    P.preparar_base_dados(caminho)
    return caminho


def test_historico_vazio(pasta_dados_temporaria):
    at = _app()
    assert any("Ainda não há importações guardadas" in i.value for i in at.info)


# 13. Filtros do histórico sem efeito nos outros separadores (e vice-versa) -------------------------------

def test_filtros_do_historico_nao_afetam_os_outros_separadores(pasta_dados_temporaria):
    at = _carregar(_app(), "dia1.xlsx", DIA1)
    _carregar(at, "dia2.xlsx", DIA2)
    antes = _estado_dos_outros_separadores(at)
    at.text_input(key="hist_pesquisa").set_value("dia1").run()
    assert list(_lista(at)["Ficheiro"]) == ["dia1.xlsx"]
    at.multiselect(key="hist_estados").set_value([H.CONCLUIDA_COM_AVISOS]).run()
    at.number_input(key="hist_pagina").set_value(1).run()
    assert _estado_dos_outros_separadores(at) == antes
    assert next(t for t in at.text_input if t.label == "Pesquisar").value == ""  # o filtro do topo não mudou


def test_filtros_do_topo_nao_afetam_o_historico(pasta_dados_temporaria):
    at = _carregar(_app(), "dia1.xlsx", DIA1)
    _carregar(at, "dia2.xlsx", DIA2)
    next(t for t in at.text_input if t.label == "Pesquisar").set_value("Ferragens").run()
    assert {m.label: m.value for m in at.metric}["Encomendas"] != "62"
    assert list(_lista(at)["Ficheiro"]) == ["dia2.xlsx", "dia1.xlsx"]
    assert _lista(at).loc[0, "Linhas"] == "62"  # totais do momento da importação, não filtrados


def test_pesquisa_sem_resultados_no_historico(pasta_dados_temporaria):
    at = _carregar(_app(), "dia1.xlsx", DIA1)
    at.text_input(key="hist_pesquisa").set_value("não existe").run()
    assert any("Nenhuma importação corresponde" in i.value for i in at.info)


# 14. Compatibilidade com as exportações Excel existentes ---------------------------------------------------

def _valores(livro: bytes) -> dict[str, list[tuple]]:
    wb = load_workbook(io.BytesIO(livro))
    return {folha.title: [tuple(c.value for c in linha) for linha in folha.iter_rows()] for folha in wb.worksheets}


def test_exportacoes_iguais_com_e_sem_historico(base_dados):
    lido, resultado, tabela, ind = processar(DIA1, "a.xlsx", DATA_REFERENCIA_EXEMPLO)
    limpos, relatorio = exportar_dados_limpos(resultado, "a.xlsx"), exportar_relatorio(resultado, tabela, ind, "a.xlsx", None)
    dados = H.preparar_dados("a.xlsx", DIA1, lido.folha, resultado.mapeamento, resultado, tabela,
                             DATA_REFERENCIA_EXEMPLO, None)
    assert H.registar_importacao(base_dados, dados).situacao == H.GUARDADA
    assert _valores(exportar_dados_limpos(resultado, "a.xlsx")) == _valores(limpos)
    assert _valores(exportar_relatorio(resultado, tabela, ind, "a.xlsx", None)) == _valores(relatorio)


def test_botoes_de_exportacao_mantem_se_com_o_historico(pasta_dados_temporaria):
    at = _carregar(_app(), "a.xlsx", DIA1)
    rotulos = [b.proto.label for b in at.get("download_button")]
    for esperado in ("Descarregar dados limpos (.xlsx)", "Descarregar relatório (.xlsx)",
                     "Descarregar estas 18 exceções (.xlsx)"):
        assert esperado in rotulos


# 11 e 12. Cópias, recuperação e base danificada na interface ------------------------------------------------

def test_criar_e_recuperar_copia_na_interface(pasta_dados_temporaria):
    at = _carregar(_app(), "dia1.xlsx", DIA1)
    at.button(key="hist_criar_copia").click().run()
    assert any("Cópia criada e verificada" in s.value for s in at.success)
    _carregar(at, "dia2.xlsx", DIA2)
    assert at.button(key="hist_recuperar").disabled  # só depois de confirmar
    at.checkbox(key="hist_confirmar_recuperacao").check().run()
    at.button(key="hist_recuperar").click().run()
    assert not at.exception and any("Histórico recuperado da cópia" in s.value for s in at.success)
    assert list(_lista(at)["Ficheiro"]) == ["dia1.xlsx"]
    assert len(list((pasta_dados_temporaria / "copias").glob("logiflow-antes-da-recuperacao-*.db"))) == 1
    # O ficheiro aberto (dia 2) não é gravado outra vez automaticamente depois da recuperação.
    at.number_input(key="prazo_padrao").set_value(2).run()
    assert H.listar_importacoes(_base())[1] == 1


def test_base_danificada_nao_impede_a_aplicacao_e_pode_ser_recuperada(pasta_dados_temporaria):
    at = _carregar(_app(), "dia1.xlsx", DIA1)
    at.button(key="hist_criar_copia").click().run()
    base = _base()
    for sufixo in ("-wal", "-shm"):
        Path(f"{base}{sufixo}").unlink(missing_ok=True)
    base.write_bytes(b"danificada" * 400)

    at = _carregar(_app(), "dia2.xlsx", DIA2)
    assert not at.exception
    assert [(m.label, m.value) for m in at.metric][0] == ("Encomendas", "62")  # o resto funciona
    assert "não pode ser lida" in at.warning[0].value
    assert any("O histórico não está disponível" in e.value for e in at.error)
    at.checkbox(key="hist_confirmar_recuperacao").check().run()
    at.button(key="hist_recuperar").click().run()
    assert any("os ficheiros originais foram guardados, sem alterações" in s.value for s in at.success)
    assert list(_lista(at)["Ficheiro"]) == ["dia1.xlsx"]  # o histórico é o da cópia, sem acrescentos
    assert any("não foi gravado outra vez" in m for m in _mensagens(at))
    at.button(key="hist_guardar_depois_de_recuperar").click().run()  # gravar o ficheiro aberto, a pedido
    assert list(_lista(at)["Ficheiro"]) == ["dia2.xlsx", "dia1.xlsx"]
    assert len(list((pasta_dados_temporaria / "copias").glob("logiflow-danificada-*.db"))) == 1


def test_pasta_de_copias_invalida_e_explicada(pasta_dados_temporaria):
    at = _carregar(_app(), "dia1.xlsx", DIA1)
    at.text_input(key="hist_pasta_copias").set_value("pasta relativa").run()
    assert any("caminho completo" in w.value for w in at.warning)
    assert not any(b.key == "hist_criar_copia" for b in at.button)


def test_pasta_de_dados_invalida_nao_impede_a_aplicacao(monkeypatch):
    monkeypatch.setenv(P.VARIAVEL_PASTA, "relativo")
    at = _carregar(_app(), "dia1.xlsx", DIA1)
    assert not at.exception and [(m.label, m.value) for m in at.metric][0] == ("Encomendas", "60")
    assert any("caminho absoluto" in w.value for w in at.warning)


# Conservação de dados e base com erros (revisão da 1.3.0) ------------------------------------------------

def _envelhecer_importacao(importacao_id: int, iniciada_em: str = "2024-01-01T08:00:00+00:00") -> None:
    with P.ligar(_base()) as ligacao, P.transacao(ligacao):
        ligacao.execute("UPDATE importacoes SET iniciada_em = ? WHERE id = ?", (iniciada_em, importacao_id))


def test_prazo_de_conservacao_pede_confirmacao_e_elimina(pasta_dados_temporaria):
    at = _carregar(_app(), "dia1.xlsx", DIA1)
    _carregar(at, "dia2.xlsx", DIA2)
    _envelhecer_importacao(1)
    at.selectbox(key="hist_prazo").set_value(12).run()
    assert any("os dados de 1 importação(ões)" in w.value for w in at.warning)
    assert at.button(key="hist_guardar_prazo").disabled  # só depois de confirmar
    at.checkbox(key="hist_confirmar_prazo").check().run()
    at.button(key="hist_guardar_prazo").click().run()
    assert not at.exception and any("Prazo de conservação" in s.value for s in at.success)
    lista = _lista(at)
    assert list(lista["Encomendas guardadas"]) == ["Sim", "Eliminadas"]
    assert list(lista["Ficheiro"]) == ["dia2.xlsx", "(eliminado)"]
    assert H.carregar_fotografia(_base(), 1).empty and R.ler_prazo(_base()) == 12
    assert any("não garante, por si só, o cumprimento do RGPD" in c.value for c in at.caption)


def test_prazo_sem_dados_afetados_nao_pede_confirmacao(pasta_dados_temporaria):
    at = _carregar(_app(), "dia1.xlsx", DIA1)
    at.selectbox(key="hist_prazo").set_value(24).run()
    assert not at.button(key="hist_guardar_prazo").disabled
    at.button(key="hist_guardar_prazo").click().run()
    assert R.ler_prazo(_base()) == 24 and len(H.carregar_fotografia(_base(), 1)) == 60


def test_eliminar_uma_importacao_na_interface(pasta_dados_temporaria):
    at = _carregar(_app(), "dia1.xlsx", DIA1)
    _carregar(at, "dia2.xlsx", DIA2)
    at.selectbox(key="hist_detalhe").set_value(1).run()
    assert at.button(key="hist_eliminar_importacao").disabled
    at.checkbox(key="hist_confirmar_eliminacao").check().run()
    at.button(key="hist_eliminar_importacao").click().run()
    assert not at.exception and any("foram eliminados" in s.value for s in at.success)
    detalhe = next(m.value for m in at.markdown if m.value.startswith("**Estado:**"))
    assert "Dados eliminados" in detalhe and "pedido manual" in detalhe
    assert not any(b.key == "hist_comparar_anterior" for b in at.button)
    assert "São precisas pelo menos duas importações concluídas" in " ".join(c.value for c in at.caption)
    assert H.carregar_fotografia(_base(), 1).empty and len(H.carregar_fotografia(_base(), 2)) == 62


def test_copias_com_dados_eliminados_assinaladas_apagadas_e_reeliminadas(pasta_dados_temporaria):
    at = _carregar(_app(), "dia1.xlsx", DIA1)
    _carregar(at, "dia2.xlsx", DIA2)
    at.button(key="hist_criar_copia").click().run()  # cópia com as duas importações
    R.eliminar_importacao(_base(), 1)
    at.run()
    tabela = next(d.value for d in at.dataframe if "Conservação" in d.value.columns)
    assert list(tabela["Conservação"]) == ["Atenção: tem dados eliminados"]
    # Recuperar essa cópia volta a eliminar os dados (não reaparecem).
    at.checkbox(key="hist_confirmar_recuperacao").check().run()
    at.button(key="hist_recuperar").click().run()
    assert any("foram eliminados de novo" in s.value for s in at.success)
    assert H.carregar_fotografia(_base(), 1).empty
    # Apagar as cópias assinaladas (com confirmação); as outras ficam.
    copias_antes = sorted(p.name for p in (pasta_dados_temporaria / "copias").glob("logiflow-*.db"))
    assert at.button(key="hist_apagar_copias").disabled
    at.checkbox(key="hist_confirmar_apagar_copias").check().run()
    at.button(key="hist_apagar_copias").click().run()
    restantes = sorted(p.name for p in (pasta_dados_temporaria / "copias").glob("logiflow-*.db"))
    assert any("Foram apagadas" in s.value for s in at.success)
    assert len(restantes) < len(copias_antes)
    for nome in restantes:  # o que ficou não tem dados eliminados
        assert not next(c for c in copias.avaliar_copias(pasta_dados_temporaria / "copias", _base())
                        if c.caminho.name == nome).assinalada


def test_conservacao_aplicada_ao_abrir_e_avisada(pasta_dados_temporaria):
    _carregar(_app(), "dia1.xlsx", DIA1)
    R.definir_prazo(_base(), 12)
    _envelhecer_importacao(1)
    at = _app()  # uma nova sessão (abrir a aplicação)
    assert any("Ao abrir a aplicação, foram eliminados os dados de 1 importação" in i.value for i in at.info)
    assert H.carregar_fotografia(_base(), 1).empty


def test_base_com_erros_fica_so_de_leitura_na_interface(pasta_dados_temporaria):
    at = _carregar(_app(), "dia1.xlsx", DIA1)
    at.button(key="hist_criar_copia").click().run()
    with P.ligar(_base()) as ligacao:
        ligacao.execute("PRAGMA wal_checkpoint(TRUNCATE)")
        paginas = ligacao.execute("PRAGMA page_count").fetchone()[0]
    dados = bytearray(_base().read_bytes())
    meio = (paginas // 2) * 4096
    dados[meio + 8:meio + 4096] = b"\x5a" * (4096 - 8)
    _base().write_bytes(bytes(dados))

    at = _carregar(_app(), "dia2.xlsx", DIA2)
    assert not at.exception and [(m.label, m.value) for m in at.metric][0] == ("Encomendas", "62")
    assert any("está só de leitura" in w.value for w in at.warning)
    assert any("tem erros e está só de leitura" in e.value for e in at.error)
    assert H.listar_importacoes(_base())[1] == 1  # nada foi gravado
    # Uma cópia feita agora fica «com erros» (não é descartada) e a recuperação da boa repõe tudo.
    at.button(key="hist_criar_copia").click().run()
    assert any("não serve para recuperar o histórico" in e.value for e in at.error)
    assert list((pasta_dados_temporaria / "copias").glob("logiflow-copia-com-erros-*.db"))
    boa = next(c for c in copias.listar_copias(pasta_dados_temporaria / "copias")
               if not c.name.startswith("logiflow-copia-com-erros"))
    at.selectbox(key="hist_copia").set_value(boa.name).run()
    at.checkbox(key="hist_confirmar_recuperacao").check().run()
    at.button(key="hist_recuperar").click().run()
    assert any("Histórico recuperado" in s.value for s in at.success)
    assert not any("só de leitura" in e.value for e in at.error)
    at.button(key="hist_guardar_depois_de_recuperar").click().run()
    assert list(_lista(at)["Ficheiro"]) == ["dia2.xlsx", "dia1.xlsx"]
