"""Testes das exportações para Excel."""

from __future__ import annotations

import datetime as dt
import io

from openpyxl import load_workbook

from logiflow.exportacao import exportar_dados_limpos, exportar_relatorio

from .conftest import CABECALHO, csv_bytes, processar

FICHEIRO = csv_bytes(
    CABECALHO + ";Data Prevista;Observações",
    "E1;Loja A;01/10/2026;Porto;Entregue;03/10/2026;Frágil",
    "E2;Loja B;31/02/2026;Braga;Pendente;05/10/2026;",          # data inválida -> erro
    "E3;Loja C;02/10/2026;Faro;em transito;04/10/2026;",        # atrasada na referência 20/10
    "E3;Loja D;03/10/2026;Évora;Pendente;30/10/2026;",           # duplicado -> erro
    # Valor malicioso, entre aspas como o Excel o gravaria num CSV.
    'E5;"=HYPERLINK(""http://exemplo.invalid"";""clique"")";04/10/2026;Viseu;Pendente;30/10/2026;=1+1',
)


def _abrir(conteudo: bytes):
    return load_workbook(io.BytesIO(conteudo))


def _linhas(folha):
    return [list(r) for r in folha.iter_rows(values_only=True)]


def test_dados_limpos_nao_perdem_nenhuma_linha():
    _, resultado, _, _ = processar(FICHEIRO)
    livro = _abrir(exportar_dados_limpos(resultado, "teste.csv"))
    assert livro.sheetnames == ["Encomendas válidas", "Linhas com erros", "Informação"]
    validas = _linhas(livro["Encomendas válidas"])[1:]
    com_erros = _linhas(livro["Linhas com erros"])[1:]
    assert len(validas) + len(com_erros) == resultado.total == 5
    assert sorted(r[0] for r in validas) == ["E1", "E5"]
    info = {r[0]: r[1] for r in _linhas(livro["Informação"]) if r[0]}
    assert info["Total confere (válidas + com erros = lidas)"] == "Sim"


def test_linhas_com_erros_mantem_valores_originais_e_motivo():
    _, resultado, _, _ = processar(FICHEIRO)
    folha = _abrir(exportar_dados_limpos(resultado, "teste.csv"))["Linhas com erros"]
    linhas = _linhas(folha)
    cabecalho = linhas[0]
    assert cabecalho[0] == "Linha no ficheiro" and cabecalho[-2:] == ["Erros", "Avisos"]
    por_linha = {r[0]: dict(zip(cabecalho, r)) for r in linhas[1:]}
    assert por_linha[3]["Data"] == "31/02/2026"        # texto original preservado
    assert "Data inválida" in por_linha[3]["Erros"]
    assert "repetido" in por_linha[4]["Erros"] and "repetido" in por_linha[5]["Erros"]


def test_dados_limpos_normalizados():
    _, resultado, _, _ = processar(FICHEIRO)
    folha = _abrir(exportar_dados_limpos(resultado, "teste.csv"))["Encomendas válidas"]
    cabecalho = [c.value for c in folha[1]]
    assert cabecalho == ["N.º Encomenda", "Cliente", "Data", "Destino", "Estado", "Data Prevista",
                         "Observações", "Avisos", "Linha no ficheiro"]
    e1 = folha[2]
    assert e1[2].value == dt.datetime(2026, 10, 1)     # data guardada como data do Excel
    assert e1[2].number_format == "DD/MM/YYYY"
    assert e1[6].value == "Frágil"                     # coluna extra preservada


def test_texto_com_formula_e_gravado_como_texto():
    """Proteção contra injeção de fórmulas: nada vindo do ficheiro é gravado como fórmula."""
    _, resultado, tabela, ind = processar(FICHEIRO)
    for conteudo in (exportar_dados_limpos(resultado, "t.csv"),
                     exportar_relatorio(resultado, tabela, ind, "t.csv")):
        livro = _abrir(conteudo)
        encontrados = 0
        for folha in livro.worksheets:
            for linha in folha.iter_rows():
                for celula in linha:
                    assert celula.data_type != "f", f"fórmula em {folha.title}!{celula.coordinate}"
                    if isinstance(celula.value, str) and celula.value.startswith("=HYPERLINK"):
                        encontrados += 1
        assert encontrados >= 1  # o valor continua lá, como texto


def test_relatorio_tem_folhas_e_valores_coerentes():
    _, resultado, tabela, ind = processar(FICHEIRO)
    livro = _abrir(exportar_relatorio(resultado, tabela, ind, "teste.csv"))
    assert livro.sheetnames == ["Resumo", "Por estado", "Atrasadas", "Problemas", "Todas as encomendas"]

    resumo = {r[0]: r[1] for r in _linhas(livro["Resumo"]) if r[0]}
    assert resumo["Total de encomendas"] == 5
    assert resumo["Entregues"] == 1
    assert resumo["Pendentes"] == 4
    # E2 tem a data da encomenda inválida, mas a data prevista (05/10) é válida: também está atrasada.
    assert resumo["Atrasadas"] == ind.atrasadas == 2
    assert resumo["Linhas com erros"] == 3
    assert resumo["Data de referência"] == dt.datetime(2026, 10, 20)

    atrasadas = _linhas(livro["Atrasadas"])
    assert len(atrasadas) - 1 == ind.atrasadas
    assert [(r[0], r[7]) for r in atrasadas[1:]] == [("E3", 16), ("E2", 15)]  # mais atrasada primeiro

    assert len(_linhas(livro["Problemas"])) - 1 == len(resultado.problemas)
    assert len(_linhas(livro["Todas as encomendas"])) - 1 == resultado.total


def test_relatorio_sem_prazos_indica_que_atrasos_nao_estao_disponiveis():
    _, resultado, tabela, ind = processar(csv_bytes(CABECALHO, "E1;A;01/10/2026;Porto;Pendente"))
    livro = _abrir(exportar_relatorio(resultado, tabela, ind, "teste.csv"))
    resumo = {r[0]: r[1] for r in _linhas(livro["Resumo"]) if r[0]}
    assert resumo["Atrasadas"] == "Não disponível"
    assert "não foram calculados" in resumo["Prazo usado para atrasos"]
    assert "Não é possível determinar atrasos" in livro["Atrasadas"]["A1"].value


def test_exportacoes_a_partir_de_excel():
    from .conftest import xlsx_bytes
    conteudo = xlsx_bytes([
        ["N.º Encomenda", "Cliente", "Data", "Destino", "Estado"],
        ["E1", "A", dt.datetime(2026, 10, 1), "Porto", "Entregue"],
        ["E2", "", dt.datetime(2026, 10, 2), "Porto", "Pendente"],
    ])
    _, resultado, tabela, ind = processar(conteudo, nome="t.xlsx")
    livro = _abrir(exportar_dados_limpos(resultado, "t.xlsx"))
    erros = _linhas(livro["Linhas com erros"])
    assert erros[1][3] == dt.datetime(2026, 10, 2)  # data original do Excel mantida como data
    assert _abrir(exportar_relatorio(resultado, tabela, ind, "t.xlsx"))["Resumo"]["A1"].value.startswith("LogiFlow")


# Relatório de desempenho -------------------------------------------------------

DESEMPENHO = csv_bytes(
    CABECALHO + ";Data Prevista;Data de Entrega",
    "A1;Armazéns Alfa;01/10/2026;Porto;Entregue;05/10/2026;04/10/2026",
    "A2;Armazéns Alfa;02/10/2026;Braga;Entregue;05/10/2026;07/10/2026",
    "A3;Armazéns Alfa;03/10/2026;Porto;Pendente;06/10/2026;",
    "B1;Bazar Beta;01/10/2026;Faro;Entregue;05/10/2026;05/10/2026",
    "B2;Bazar Beta;02/10/2026;Évora;Pendente;30/10/2026;",
    "B2;Bazar Beta;02/10/2026;Évora;Pendente;30/10/2026;",          # duplicado: fica de fora
    '"C1";"=HYPERLINK(""http://exemplo.invalid"";""x"")";01/10/2026;"=1+1";Entregue;05/10/2026;06/10/2026',
)


def _valores(livro) -> list[str]:
    return [str(c.value) for ws in livro.worksheets for linha in ws.iter_rows() for c in linha if c.value is not None]


def test_relatorio_de_desempenho_de_todos_os_clientes():
    from logiflow import desempenho as D
    from logiflow.exportacao import exportar_desempenho

    _, resultado, tabela, _ = processar(DESEMPENHO)
    livro = _abrir(exportar_desempenho(resultado, tabela, "teste.csv", None, D.SEMANA))
    assert livro.sheetnames == ["Resumo", "Por período", "Por cliente", "Por destino",
                                "Pendentes por antiguidade", "Entregas com atraso"]
    resumo = {r[0]: r[1:] for r in _linhas(livro["Resumo"]) if r[0]}
    r = D.resumir(tabela)
    assert resumo["Âmbito"][0] == "Todos os clientes"
    assert "2 linhas com erros ficaram de fora" in resumo["Base de cálculo"][0]
    assert resumo["Taxa de entregas no prazo"][0] == r.taxa_no_prazo == 0.5
    assert resumo["Taxa de entregas no prazo"][1] == "2 de 4 entregas avaliadas — poucas entregas " \
                                                      "(menos de 5): use como indicação"
    folha = livro["Resumo"]
    celula_taxa = next(c for c in folha["B"] if c.value == r.taxa_no_prazo)
    assert celula_taxa.number_format == "0.0%"
    clientes = [linha[0] for linha in _linhas(livro["Por cliente"])[1:]]
    assert set(clientes) == {"Armazéns Alfa", "Bazar Beta", '=HYPERLINK("http://exemplo.invalid";"x")'}
    atrasos = _linhas(livro["Entregas com atraso"])[1:]
    assert [(a[0], a[6]) for a in atrasos] == [("A2", 2), ("C1", 1)]


def test_relatorio_de_desempenho_de_um_cliente_nao_tem_dados_de_outros():
    """O relatório por cliente pode ser enviado ao cliente: não contém dados de mais ninguém."""
    from logiflow.exportacao import exportar_desempenho

    _, resultado, tabela, _ = processar(DESEMPENHO)
    livro = _abrir(exportar_desempenho(resultado, tabela, "teste.csv", cliente="ARMAZENS ALFA"))
    assert "Por cliente" not in livro.sheetnames
    valores = _valores(livro)
    assert any("Cliente «ARMAZENS ALFA»" in v for v in valores)
    assert "3 encomendas sem erros" in valores
    for alheio in ("Bazar Beta", "B1", "B2", "Faro", "Évora", "HYPERLINK", "C1"):
        assert not any(alheio in v for v in valores), alheio
    resumo = {r[0]: r[1:] for r in _linhas(livro["Resumo"]) if r[0]}
    assert resumo["Taxa de entregas no prazo"][0] == 0.5      # A1 no prazo, A2 com atraso
    assert resumo["Pendentes atrasadas"][0] == 1              # A3


def test_relatorio_de_desempenho_nao_grava_formulas():
    from logiflow.exportacao import exportar_desempenho

    _, resultado, tabela, _ = processar(DESEMPENHO)
    livro = _abrir(exportar_desempenho(resultado, tabela, "teste.csv"))
    for folha in livro.worksheets:
        for linha in folha.iter_rows():
            for celula in linha:
                assert celula.data_type != "f", f"fórmula em {folha.title}!{celula.coordinate}"
    assert any(v.startswith("=HYPERLINK") for v in _valores(livro))


def test_relatorio_de_desempenho_sem_datas_indica_nao_disponivel():
    from logiflow import desempenho as D
    from logiflow.exportacao import exportar_desempenho

    _, resultado, tabela, _ = processar(csv_bytes(CABECALHO, "E1;A;01/10/2026;Porto;Entregue",
                                                  "E2;A;01/10/2026;Porto;Pendente"))
    livro = _abrir(exportar_desempenho(resultado, tabela, "teste.csv", periodo=D.MES))
    resumo = {r[0]: r[1:] for r in _linhas(livro["Resumo"]) if r[0]}
    assert resumo["Taxa de entregas no prazo"][0] == "Não disponível"
    assert resumo["Tempo médio de entrega (dias)"][0] == "Não disponível"
    assert resumo["Pendentes atrasadas"][0] == "Não disponível"
    assert _linhas(livro["Por período"])[0][0] == "Mês"


def test_relatorio_de_desempenho_concordancia_no_singular():
    from logiflow.exportacao import exportar_desempenho

    _, resultado, tabela, _ = processar(csv_bytes(CABECALHO, "E1;A;01/10/2026;Porto;Entregue",
                                                  "E2;;01/10/2026;Porto;Entregue"))
    resumo = {r[0]: r[1] for r in _linhas(_abrir(exportar_desempenho(resultado, tabela, "t.csv"))["Resumo"]) if r[0]}
    assert resumo["Base de cálculo"] == "1 encomenda sem erros (1 linha com erros ficou de fora)"


def test_relatorio_de_desempenho_pronto_a_imprimir():
    from logiflow.exportacao import exportar_desempenho

    _, resultado, tabela, _ = processar(DESEMPENHO)
    for folha in _abrir(exportar_desempenho(resultado, tabela, "t.csv")).worksheets:
        assert folha.page_setup.orientation == "landscape"
        assert folha.sheet_properties.pageSetUpPr.fitToPage and folha.page_setup.fitToWidth == 1


# Cabeçalhos vindos do ficheiro (correção de segurança na 1.2) ------------------------

def test_nome_de_coluna_com_formula_e_gravado_como_texto():
    """Uma coluna extra cujo nome é uma fórmula: o cabeçalho exportado tem de ser texto."""
    conteudo = csv_bytes(CABECALHO + ';"=HYPERLINK(""http://exemplo.invalid"";""x"")"',
                         "E1;A;01/10/2026;Porto;Pendente;valor", "E2;;01/10/2026;Porto;Pendente;valor")
    _, resultado, _, _ = processar(conteudo)
    livro = _abrir(exportar_dados_limpos(resultado, "t.csv"))
    cabecalhos = []
    for folha in livro.worksheets:
        for linha in folha.iter_rows():
            for celula in linha:
                assert celula.data_type != "f", f"fórmula em {folha.title}!{celula.coordinate}"
        cabecalhos += [c.value for c in folha[1]]
    assert cabecalhos.count('=HYPERLINK("http://exemplo.invalid";"x")') == 2  # válidas e com erros


# Exceções -------------------------------------------------------------------------------

EXCECOES = csv_bytes(
    CABECALHO + ";Data Prevista;Data de Entrega",
    "A1;Armazéns Alfa;01/10/2026;Porto;Pendente;05/10/2026;",        # Alta
    "A2;Armazéns Alfa;18/10/2026;Braga;Pendente;21/10/2026;",        # Média (vence amanhã)
    "B1;Bazar Beta;01/10/2026;Faro;Em trânsito;05/10/2026;",         # Alta
    "B2;Bazar Beta;31/02/2026;Évora;Pendente;30/10/2026;",           # Baixa (data inválida)
    '"C1";"=HYPERLINK(""http://exemplo.invalid"";""x"")";01/10/2026;"=1+1";Pendente;05/10/2026;',  # Alta
)


def _excecoes_exemplo():
    from logiflow import excecoes as E

    _, resultado, tabela, _ = processar(EXCECOES)
    return E, E.identificar_excecoes(tabela, resultado.problemas, resultado.data_referencia)


def test_exportacao_de_excecoes_so_com_as_filtradas_e_pela_mesma_ordem():
    from logiflow.exportacao import exportar_excecoes

    E, todas = _excecoes_exemplo()
    filtradas = E.filtrar(todas, prioridades=[E.ALTA])
    livro = _abrir(exportar_excecoes(filtradas, "t.csv", dt.date(2026, 10, 20), [("Prioridades", "Alta")]))
    assert livro.sheetnames == ["Exceções", "Informação"]
    linhas = _linhas(livro["Exceções"])
    assert linhas[0][:5] == ["Prioridade", "N.º Encomenda", "Cliente", "Categorias", "Motivos"]
    assert linhas[0][-1] == "Linha no ficheiro"
    assert [r[1] for r in linhas[1:]] == list(filtradas[E.NUMERO]) and len(linhas) - 1 == 3
    assert all(r[0] == "Alta" and r[4] for r in linhas[1:])  # prioridade e motivos incluídos
    info = {r[0]: r[1] for r in _linhas(livro["Informação"]) if r[0]}
    assert info["Encomendas exportadas"] == 3 and info["Prioridades"] == "Alta"
    assert info["Média"] == 0 and info["Baixa"] == 0  # totais só das linhas exportadas


def test_exportacao_de_excecoes_sem_formulas():
    from logiflow.exportacao import exportar_excecoes

    E, todas = _excecoes_exemplo()
    livro = _abrir(exportar_excecoes(todas, "t.csv", dt.date(2026, 10, 20), [("Pesquisa", "=1+2")]))
    for folha in livro.worksheets:
        for linha in folha.iter_rows():
            for celula in linha:
                assert celula.data_type != "f", f"fórmula em {folha.title}!{celula.coordinate}"
    valores = _valores(livro)
    assert '=HYPERLINK("http://exemplo.invalid";"x")' in valores and "=1+1" in valores and "=1+2" in valores


def test_exportacao_de_excecoes_de_um_cliente_nao_tem_dados_de_outros():
    from logiflow.exportacao import exportar_excecoes

    E, todas = _excecoes_exemplo()
    so_alfa = E.filtrar(todas, cliente="armazens alfa")
    livro = _abrir(exportar_excecoes(so_alfa, "t.csv", dt.date(2026, 10, 20), [("Cliente", "Armazéns Alfa")]))
    valores = _valores(livro)
    for alheio in ("Bazar Beta", "B1", "B2", "Faro", "Évora", "HYPERLINK", "C1", "=1+1"):
        assert not any(alheio in v for v in valores), alheio
    info = {r[0]: r[1] for r in _linhas(livro["Informação"]) if r[0]}
    assert info["Encomendas exportadas"] == 2 and info["Alta"] == 1 and info["Média"] == 1


def test_exportacao_sem_excecoes_gera_ficheiro_valido_e_explicito():
    from logiflow import excecoes as E
    from logiflow.exportacao import exportar_excecoes

    _, todas = _excecoes_exemplo()
    vazio = E.filtrar(todas, pesquisa="nao existe")
    livro = _abrir(exportar_excecoes(vazio, "t.csv", dt.date(2026, 10, 20)))
    linhas = _linhas(livro["Exceções"])
    assert linhas[0][0] == "Prioridade" and linhas[2][0] == "Não há exceções com os filtros aplicados."
    info = {r[0]: r[1] for r in _linhas(livro["Informação"]) if r[0]}
    assert info["Encomendas exportadas"] == 0 and info["Alta"] == 0
    assert livro["Exceções"].page_setup.orientation == "landscape"
