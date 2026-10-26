"""Leitura segura de ficheiros CSV e Excel (.xlsx).

O conteúdo dos ficheiros é tratado sempre como dados, nunca como código:

* Os ficheiros Excel são lidos com o openpyxl em modo "só valores": as fórmulas
  não são calculadas (usa-se o último valor guardado pelo Excel) e as macros
  nunca são executadas. Com o pacote ``defusedxml`` instalado, o openpyxl
  recusa também XML malicioso.
* Antes de abrir um .xlsx (que é um ficheiro ZIP) verifica-se o tamanho
  descomprimido, para recusar "bombas ZIP".
* Há limites de tamanho, de linhas e de colunas, com mensagens claras.
* Nenhum valor é descartado sem aviso: linhas totalmente vazias e linhas
  antes do cabeçalho são contadas e comunicadas ao utilizador.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import re
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from pathlib import PurePath

import pandas as pd

from .campos import campo_da_coluna

TAMANHO_MAXIMO_MB = 20
TAMANHO_MAXIMO = TAMANHO_MAXIMO_MB * 1024 * 1024
TAMANHO_DESCOMPRIMIDO_MAXIMO = 300 * 1024 * 1024
MAX_LINHAS = 200_000
MAX_COLUNAS = 200
MAX_CARACTERES_CELULA = 32_767  # limite do próprio Excel
LINHAS_PROCURA_CABECALHO = 20

EXTENSOES_ACEITES = (".csv", ".xlsx")
_SEPARADORES = (";", ",", "\t", "|")
_NOMES_SEPARADOR = {";": "ponto e vírgula (;)", ",": "vírgula (,)", "\t": "tabulação", "|": "barra vertical (|)"}

# Caracteres de controlo invisíveis que o formato Excel não aceita.
_CARACTERES_CONTROLO = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

_ASSINATURA_ZIP = b"PK\x03\x04"
_ASSINATURA_OLE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"  # .xls antigo ou ficheiro com palavra-passe


class ErroImportacao(Exception):
    """Erro com uma mensagem pronta a mostrar ao utilizador."""


@dataclass
class FicheiroLido:
    nome: str
    formato: str                      # "CSV" ou "Excel"
    dados: pd.DataFrame               # valores tal como lidos (texto ou datas do Excel)
    linhas: list[int]                 # n.º da linha no ficheiro original, para cada registo
    linha_cabecalho: int
    folha: str | None = None
    codificacao: str | None = None
    separador: str | None = None
    linhas_vazias_ignoradas: int = 0
    avisos: list[str] = field(default_factory=list)

    @property
    def colunas(self) -> list[str]:
        return list(self.dados.columns)

    @property
    def total_linhas(self) -> int:
        return len(self.dados)


# Funções públicas ---------------------------------------------------------

def listar_folhas(conteudo: bytes) -> list[str]:
    """Nomes das folhas de um ficheiro .xlsx."""
    livro = _abrir_excel(conteudo)
    try:
        return list(livro.sheetnames)
    finally:
        livro.close()


def ler_ficheiro(nome: str, conteudo: bytes, folha: str | None = None) -> FicheiroLido:
    """Lê um ficheiro CSV ou .xlsx e devolve os registos encontrados.

    Lança ``ErroImportacao`` com uma explicação se o ficheiro não puder ser lido.
    """
    extensao = PurePath(nome).suffix.lower()
    if extensao == ".xls":
        raise ErroImportacao(
            "O formato .xls (Excel 97-2003) não é suportado. Abra o ficheiro no Excel "
            "e use «Guardar como» → «Livro do Excel (*.xlsx)»."
        )
    if extensao not in EXTENSOES_ACEITES:
        raise ErroImportacao(
            f"Tipo de ficheiro não suportado ({extensao or 'sem extensão'}). "
            "Use um ficheiro CSV (.csv) ou Excel (.xlsx)."
        )
    if not conteudo:
        raise ErroImportacao("O ficheiro está vazio (0 bytes).")
    if len(conteudo) > TAMANHO_MAXIMO:
        raise ErroImportacao(
            f"O ficheiro tem {len(conteudo) / 1024 / 1024:.1f} MB e o limite é "
            f"{TAMANHO_MAXIMO_MB} MB. Divida-o em ficheiros mais pequenos."
        )

    if extensao == ".xlsx":
        linhas_brutas, folha_lida, avisos = _ler_linhas_excel(conteudo, folha)
        resultado = _construir(nome, "Excel", linhas_brutas, avisos)
        resultado.folha = folha_lida
    else:
        linhas_brutas, codificacao, separador, avisos = _ler_linhas_csv(conteudo)
        resultado = _construir(nome, "CSV", linhas_brutas, avisos)
        resultado.codificacao = codificacao
        resultado.separador = _NOMES_SEPARADOR.get(separador, separador)
    return resultado


# Excel --------------------------------------------------------------------

def _abrir_excel(conteudo: bytes):
    if conteudo.startswith(_ASSINATURA_OLE):
        raise ErroImportacao(
            "Este ficheiro está no formato Excel antigo (.xls) ou está protegido por "
            "palavra-passe. Abra-o no Excel, retire a proteção se existir e guarde-o "
            "como «Livro do Excel (*.xlsx)»."
        )
    if not conteudo.startswith(_ASSINATURA_ZIP):
        raise ErroImportacao(
            "O ficheiro não é um Excel (.xlsx) válido. Se for um CSV, altere a extensão "
            "para .csv; caso contrário, volte a guardá-lo no Excel como .xlsx."
        )
    try:
        with zipfile.ZipFile(io.BytesIO(conteudo)) as arquivo:
            descomprimido = sum(info.file_size for info in arquivo.infolist())
    except zipfile.BadZipFile:
        raise ErroImportacao("O ficheiro Excel está danificado e não pode ser aberto.") from None
    if descomprimido > TAMANHO_DESCOMPRIMIDO_MAXIMO:
        raise ErroImportacao(
            "O conteúdo do ficheiro Excel é demasiado grande depois de descomprimido. "
            "Por segurança, não foi aberto."
        )

    from openpyxl import load_workbook  # importação local: só é preciso para .xlsx

    try:
        # read_only: leitura eficiente; data_only: usa valores, nunca fórmulas.
        return load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    except Exception as erro:  # o openpyxl lança vários tipos de erro diferentes
        raise ErroImportacao(
            "Não foi possível ler o ficheiro Excel. Confirme que abre no Excel e "
            f"volte a guardá-lo como .xlsx. (Detalhe técnico: {type(erro).__name__})"
        ) from None


def _ler_linhas_excel(conteudo: bytes, folha: str | None) -> tuple[list[list[object]], str, list[str]]:
    livro = _abrir_excel(conteudo)
    try:
        if not livro.sheetnames:
            raise ErroImportacao("O ficheiro Excel não tem folhas.")
        if folha is None:
            folha = livro.sheetnames[0]
        elif folha not in livro.sheetnames:
            raise ErroImportacao(f"A folha «{folha}» não existe neste ficheiro.")
        avisos: list[str] = []
        if len(livro.sheetnames) > 1:
            avisos.append(
                f"O ficheiro tem {len(livro.sheetnames)} folhas; foi lida a folha «{folha}». "
                "Pode escolher outra na barra lateral."
            )
        linhas: list[list[object]] = []
        for linha in livro[folha].iter_rows(values_only=True):
            if len(linhas) > MAX_LINHAS:
                raise ErroImportacao(_mensagem_linhas_a_mais())
            linhas.append([_valor_excel(v) for v in linha])
        return linhas, folha, avisos
    finally:
        livro.close()


def _valor_excel(valor: object) -> object:
    """Normaliza um valor de célula: datas ficam datas; o resto passa a texto."""
    if valor is None:
        return ""
    if isinstance(valor, (dt.datetime, dt.date)):
        return valor
    if isinstance(valor, bool):
        return "VERDADEIRO" if valor else "FALSO"
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor)


# CSV ----------------------------------------------------------------------

def _descodificar(conteudo: bytes) -> tuple[str, str]:
    if conteudo.startswith((b"\xff\xfe", b"\xfe\xff")):
        try:
            return conteudo.decode("utf-16"), "UTF-16"
        except UnicodeDecodeError:
            pass
    for codificacao, nome in (("utf-8-sig", "UTF-8"), ("cp1252", "Windows-1252")):
        try:
            return conteudo.decode(codificacao), nome
        except UnicodeDecodeError:
            continue
    return conteudo.decode("latin-1"), "ISO-8859-1"


def _detetar_separador(texto: str) -> str:
    """Escolhe o separador que produz um número de colunas mais consistente."""
    amostra = [l for l in texto.splitlines()[:30] if l.strip()]
    melhor, melhor_pontuacao = ";", (0, 0)
    for separador in _SEPARADORES:
        try:
            contagens = [len(l) for l in csv.reader(amostra, delimiter=separador)]
        except csv.Error:
            continue
        if not contagens:
            continue
        colunas, frequencia = Counter(contagens).most_common(1)[0]
        if colunas < 2:
            continue
        pontuacao = (frequencia, colunas)
        if pontuacao > melhor_pontuacao:
            melhor, melhor_pontuacao = separador, pontuacao
    return melhor


def _ler_linhas_csv(conteudo: bytes) -> tuple[list[list[object]], str, str, list[str]]:
    if conteudo.startswith(_ASSINATURA_ZIP):
        raise ErroImportacao(
            "Este ficheiro parece ser um Excel (.xlsx) com a extensão .csv. "
            "Altere a extensão para .xlsx e importe-o novamente."
        )
    texto, codificacao = _descodificar(conteudo)
    if "\x00" in texto:
        raise ErroImportacao("O ficheiro não parece ser um CSV de texto (contém dados binários).")
    avisos: list[str] = []
    if codificacao == "ISO-8859-1":
        avisos.append(
            "A codificação do ficheiro não foi identificada com certeza; alguns acentos "
            "podem aparecer trocados. Se possível, guarde o CSV como «CSV UTF-8»."
        )
    separador = _detetar_separador(texto)
    linhas: list[list[object]] = []
    # strict=True: umas aspas por fechar fariam o resto do ficheiro "desaparecer"
    # dentro de um único campo. Assim, o erro é comunicado em vez de perder linhas.
    leitor = csv.reader(io.StringIO(texto, newline=""), delimiter=separador, strict=True)
    try:
        for linha in leitor:
            if len(linhas) > MAX_LINHAS:
                raise ErroImportacao(_mensagem_linhas_a_mais())
            linhas.append(list(linha))
    except csv.Error as erro:
        raise ErroImportacao(
            f"O CSV tem um erro de formato perto da linha {leitor.line_num} ({erro}). "
            "Verifique se há aspas (\") por fechar nessa zona do ficheiro."
        ) from None
    return linhas, codificacao, separador, avisos


# Construção do resultado --------------------------------------------------

def _mensagem_linhas_a_mais() -> str:
    return (
        f"O ficheiro tem mais de {MAX_LINHAS:,} linhas, que é o limite desta versão. "
        "Divida-o em ficheiros mais pequenos (por exemplo, por mês)."
    ).replace(",", " ")


def _vazio(valor: object) -> bool:
    return isinstance(valor, str) and not valor.strip()


def _encontrar_cabecalho(linhas: list[list[object]]) -> int:
    """Índice da linha de cabeçalho.

    É a primeira linha (entre as primeiras 20 não vazias) com pelo menos dois
    nomes de coluna reconhecidos. Se nenhuma tiver, usa-se a primeira linha não vazia.
    """
    primeira_nao_vazia = None
    vistas = 0
    for indice, linha in enumerate(linhas):
        if all(_vazio(v) for v in linha):
            continue
        if primeira_nao_vazia is None:
            primeira_nao_vazia = indice
        reconhecidos = {campo_da_coluna(str(v)) for v in linha if not _vazio(v)} - {None}
        if len(reconhecidos) >= 2:
            return indice
        vistas += 1
        if vistas >= LINHAS_PROCURA_CABECALHO:
            break
    if primeira_nao_vazia is None:
        raise ErroImportacao("O ficheiro não contém dados (todas as linhas estão vazias).")
    return primeira_nao_vazia


def _nome_coluna_excel(indice: int) -> str:
    """0 -> A, 25 -> Z, 26 -> AA."""
    nome = ""
    indice += 1
    while indice:
        indice, resto = divmod(indice - 1, 26)
        nome = chr(65 + resto) + nome
    return nome


def _limpar_texto(valor: object, contador: list[int], linha: int) -> object:
    if isinstance(valor, str):
        limpo, n = _CARACTERES_CONTROLO.subn("", valor)
        contador[0] += n
        if len(limpo) > MAX_CARACTERES_CELULA:
            raise ErroImportacao(
                f"A linha {linha} tem um valor com mais de {MAX_CARACTERES_CELULA} caracteres, "
                "o máximo que o Excel aceita numa célula. Verifique se o ficheiro está correto."
            )
        return limpo.strip()
    return valor


def _construir(nome: str, formato: str, linhas: list[list[object]], avisos: list[str]) -> FicheiroLido:
    if not linhas:
        raise ErroImportacao("O ficheiro não contém dados.")

    indice_cabecalho = _encontrar_cabecalho(linhas)
    antes_cabecalho = sum(
        1 for linha in linhas[:indice_cabecalho] if not all(_vazio(v) for v in linha)
    )
    if antes_cabecalho:
        avisos.append(
            f"O cabeçalho foi encontrado na linha {indice_cabecalho + 1}; "
            f"{antes_cabecalho} linha(s) anterior(es) com texto (por exemplo, um título) "
            "não foram tratadas como encomendas."
        )

    corpo = linhas[indice_cabecalho + 1:]
    largura = max(len(l) for l in [linhas[indice_cabecalho], *corpo])
    if largura > MAX_COLUNAS:
        raise ErroImportacao(
            f"O ficheiro tem {largura} colunas; o limite é {MAX_COLUNAS}. "
            "Confirme que o separador do CSV está correto ou retire colunas desnecessárias."
        )

    def completar(linha: list[object]) -> list[object]:
        return list(linha) + [""] * (largura - len(linha))

    cabecalho = completar(linhas[indice_cabecalho])
    registos: list[list[object]] = []
    numeros_linha: list[int] = []
    vazias = 0
    controlo = [0]
    for posicao, linha in enumerate(corpo, start=indice_cabecalho + 2):
        valores = [_limpar_texto(v, controlo, posicao) for v in completar(linha)]
        if all(_vazio(v) for v in valores):
            vazias += 1
            continue
        registos.append(valores)
        numeros_linha.append(posicao)

    # Nomes de coluna: sem nome -> "Coluna sem nome (C)"; repetidos -> "Nome (2)".
    nomes: list[str] = []
    manter: list[int] = []
    for indice, valor in enumerate(cabecalho):
        texto = valor.strftime("%d/%m/%Y") if isinstance(valor, (dt.date, dt.datetime)) else str(valor)
        texto = _CARACTERES_CONTROLO.sub("", texto).strip()
        tem_valores = any(not _vazio(r[indice]) for r in registos)
        if not texto:
            if not tem_valores:
                continue  # coluna sem cabeçalho nem dados (comum no fim das folhas Excel)
            texto = f"Coluna sem nome ({_nome_coluna_excel(indice)})"
            avisos.append(f"A coluna {_nome_coluna_excel(indice)} tem dados mas não tem cabeçalho; foi chamada «{texto}».")
        base, n = texto, 2
        while texto in nomes:
            texto = f"{base} ({n})"
            n += 1
        if texto != base:
            avisos.append(f"Há mais do que uma coluna chamada «{base}»; a repetida passou a chamar-se «{texto}».")
        nomes.append(texto)
        manter.append(indice)

    if not registos:
        raise ErroImportacao(
            "O ficheiro só tem o cabeçalho e não contém nenhuma encomenda."
        )

    dados = pd.DataFrame([[r[i] for i in manter] for r in registos], columns=nomes, dtype=object)

    if vazias:
        avisos.append(f"{vazias} linha(s) completamente vazia(s) foram ignoradas (não continham dados).")
    if controlo[0]:
        avisos.append(f"Foram removidos {controlo[0]} caracteres de controlo invisíveis do conteúdo.")

    return FicheiroLido(
        nome=nome,
        formato=formato,
        dados=dados,
        linhas=numeros_linha,
        linha_cabecalho=indice_cabecalho + 1,
        linhas_vazias_ignoradas=vazias,
        avisos=avisos,
    )
