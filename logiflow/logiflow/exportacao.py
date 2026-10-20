"""Exportação para Excel: dados limpos, relatório, relatório de desempenho e exceções.

Garantias:

* **Nada se perde.** O ficheiro de dados limpos tem as linhas válidas numa
  folha e as linhas com erros noutra, com os valores originais e o motivo.
  A folha «Informação» mostra que as contas batem certo.
* **Sem fórmulas vindas dos dados.** Um texto como ``=HYPERLINK(...)`` vindo do
  ficheiro de origem é gravado como texto simples, nunca como fórmula.
"""

from __future__ import annotations

import datetime as dt
import io
import math
from collections.abc import Iterable, Sequence

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from . import __version__
from . import campos as C
from . import desempenho as D
from . import excecoes as E
from .datas import formatar_data
from .indicadores import (
    COL_DIAS_ATRASO,
    COL_ORIGEM_PRAZO,
    COL_PRAZO,
    COL_SITUACAO,
    SIT_ATRASADA,
    SITUACOES,
    Indicadores,
)
from .validacao import (
    AVISO,
    COL_AVISOS,
    COL_ERROS,
    COL_LINHA,
    COL_TEM_ERRO,
    ERRO,
    ResultadoValidacao,
)

FORMATO_DATA = "DD/MM/YYYY"
FORMATO_PERCENTAGEM = "0.0%"
FORMATO_DIAS = "0.0"
_COR_CABECALHO = "1F4E79"
_FONTE_CABECALHO = Font(bold=True, color="FFFFFF")
_FUNDO_CABECALHO = PatternFill("solid", fgColor=_COR_CABECALHO)
_FONTE_TITULO = Font(bold=True, size=14, color=_COR_CABECALHO)
_FONTE_SECCAO = Font(bold=True, size=12, color=_COR_CABECALHO)
_LIMITE = Border(bottom=Side(style="thin", color="BFBFBF"))

ROTULO_LINHA = "Linha no ficheiro"


# Escrita de células ---------------------------------------------------------

def _valor_celula(valor: object) -> object:
    """Converte valores do pandas/Python em valores que o Excel aceita."""
    if valor is None:
        return None
    if valor is pd.NaT or (isinstance(valor, float) and math.isnan(valor)):
        return None
    try:
        if pd.isna(valor):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(valor, pd.Timestamp):
        return valor.to_pydatetime().date()
    if isinstance(valor, dt.datetime):
        return valor.date()
    if isinstance(valor, bool):
        return "Sim" if valor else "Não"
    if hasattr(valor, "item"):  # números do numpy
        return valor.item()
    return valor


def _escrever(ws: Worksheet, linha: int, coluna: int, valor: object, negrito: bool = False):
    celula = ws.cell(row=linha, column=coluna, value=_valor_celula(valor))
    # O openpyxl trata qualquer texto começado por "=" como fórmula. Os dados
    # vêm de ficheiros externos, por isso forçamos sempre o tipo "texto".
    if celula.data_type == "f":
        celula.data_type = "s"
    if isinstance(celula.value, dt.date):
        celula.number_format = FORMATO_DATA
    if negrito:
        celula.font = Font(bold=True)
    return celula


def _escrever_tabela(ws: Worksheet, cabecalho: Sequence[str], linhas: Iterable[Sequence[object]],
                     linha_inicial: int = 1, filtro: bool = True) -> int:
    """Escreve uma tabela formatada e devolve a última linha usada."""
    for j, titulo in enumerate(cabecalho, start=1):
        celula = ws.cell(row=linha_inicial, column=j, value=titulo)
        # Alguns cabeçalhos são nomes de colunas do ficheiro importado: também nunca como fórmula.
        if celula.data_type == "f":
            celula.data_type = "s"
        celula.font = _FONTE_CABECALHO
        celula.fill = _FUNDO_CABECALHO
        celula.alignment = Alignment(vertical="center", wrap_text=True)
    ultima = linha_inicial
    larguras = [len(str(t)) for t in cabecalho]
    for i, valores in enumerate(linhas, start=linha_inicial + 1):
        for j, valor in enumerate(valores, start=1):
            celula = _escrever(ws, i, j, valor)
            if i - linha_inicial <= 300:
                texto = formatar_data(celula.value) if isinstance(celula.value, dt.date) else str(celula.value or "")
                larguras[j - 1] = max(larguras[j - 1], len(texto))
        ultima = i
    for j, largura in enumerate(larguras, start=1):
        letra = get_column_letter(j)
        ws.column_dimensions[letra].width = max(10, min(60, largura + 2))
    if filtro and cabecalho:
        ws.auto_filter.ref = f"A{linha_inicial}:{get_column_letter(len(cabecalho))}{max(ultima, linha_inicial)}"
        ws.freeze_panes = ws.cell(row=linha_inicial + 1, column=1)
    return ultima


def _escrever_pares(ws: Worksheet, linha: int, pares: Sequence[tuple[str, object]]) -> int:
    for rotulo, valor in pares:
        _escrever(ws, linha, 1, rotulo, negrito=True)
        _escrever(ws, linha, 2, valor)
        ws.cell(row=linha, column=1).border = _LIMITE
        ws.cell(row=linha, column=2).border = _LIMITE
        linha += 1
    return linha


def _formatar_colunas(ws: Worksheet, primeira: int, ultima: int, formatos: dict[int, str]) -> None:
    """Aplica formatos numéricos (por exemplo, percentagem) a colunas de uma tabela."""
    for coluna, formato in formatos.items():
        for linha in range(primeira, ultima + 1):
            ws.cell(row=linha, column=coluna).number_format = formato


def _escrever_notas(ws: Worksheet, linha: int, titulo: str, notas: Sequence[str], largura: int = 4) -> int:
    _escrever(ws, linha, 1, titulo).font = _FONTE_SECCAO
    for k, nota in enumerate(notas, start=linha + 1):
        _escrever(ws, k, 1, f"• {nota}").alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=k, start_column=1, end_row=k, end_column=largura)
        ws.row_dimensions[k].height = 30
    return linha + len(notas) + 1


def _gravar(livro: Workbook) -> bytes:
    memoria = io.BytesIO()
    livro.save(memoria)
    return memoria.getvalue()


def _agora() -> str:
    return dt.datetime.now().strftime("%d/%m/%Y %H:%M")


def _linhas(df: pd.DataFrame, colunas: Sequence[str], mascara: pd.Series | None = None) -> list[tuple]:
    """Linhas de um DataFrame como tuplos (muito mais rápido do que ler célula a célula)."""
    if mascara is not None:
        df = df[mascara]
    return list(zip(*(df[c].tolist() for c in colunas))) if colunas else [() for _ in range(len(df))]


# Dados limpos ---------------------------------------------------------------

def _campos_exportados(resultado: ResultadoValidacao) -> list[str]:
    return [c.nome for c in C.CAMPOS if c.nome in resultado.campos_presentes]


def exportar_dados_limpos(resultado: ResultadoValidacao, nome_origem: str) -> bytes:
    """Excel com as linhas válidas (normalizadas) e as linhas com erros (originais)."""
    tabela = resultado.tabela
    extras = resultado.extras
    campos = _campos_exportados(resultado)
    erro = tabela[COL_TEM_ERRO]
    n_validas, n_com_erro = int((~erro).sum()), int(erro.sum())

    livro = Workbook()

    # 1. Linhas válidas, com valores normalizados (datas como datas, estado uniformizado).
    ws = livro.active
    ws.title = "Encomendas válidas"
    cabecalho = [C.rotulo(c) for c in campos] + list(extras.columns) + ["Avisos", ROTULO_LINHA]
    _escrever_tabela(ws, cabecalho, (
        a + b + c for a, b, c in zip(
            _linhas(tabela, campos, ~erro),
            _linhas(extras, list(extras.columns), ~erro),
            _linhas(tabela, [COL_AVISOS, COL_LINHA], ~erro),
        )
    ))

    # 2. Linhas com erros, exatamente como estavam no ficheiro de origem.
    ws = livro.create_sheet("Linhas com erros")
    originais = resultado.originais
    cabecalho = [ROTULO_LINHA] + list(originais.columns) + ["Erros", "Avisos"]
    _escrever_tabela(ws, cabecalho, (
        a + b + c for a, b, c in zip(
            _linhas(tabela, [COL_LINHA], erro),
            _linhas(originais, list(originais.columns), erro),
            _linhas(tabela, [COL_ERROS, COL_AVISOS], erro),
        )
    ))

    # 3. Informação e verificação de totais.
    ws = livro.create_sheet("Informação")
    _escrever(ws, 1, 1, "LogiFlow — Dados limpos").font = _FONTE_TITULO
    _escrever_pares(ws, 3, [
        ("Ficheiro de origem", nome_origem),
        ("Exportado em", _agora()),
        ("Encomendas lidas", resultado.total),
        ("Encomendas válidas", n_validas),
        ("Linhas com erros", n_com_erro),
        ("Total confere (válidas + com erros = lidas)",
         "Sim" if n_validas + n_com_erro == resultado.total else "NÃO — verifique"),
    ])
    notas = [
        "As linhas com erros não foram apagadas: estão na folha «Linhas com erros», com os valores "
        "originais e a descrição do problema. Corrija-as no ficheiro de origem e importe-o de novo.",
        "Na folha «Encomendas válidas», as datas foram convertidas para datas do Excel, os espaços "
        "nas pontas foram retirados e os estados reconhecidos foram uniformizados "
        "(por exemplo, «em transito» passa a «Em trânsito»).",
        "As linhas válidas podem ter avisos (coluna «Avisos») que convém confirmar.",
    ]
    for k, nota in enumerate(notas, start=11):
        _escrever(ws, k, 1, nota).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=k, start_column=1, end_row=k, end_column=2)
        ws.row_dimensions[k].height = 45
    ws.column_dimensions["A"].width = 45
    ws.column_dimensions["B"].width = 60
    return _gravar(livro)


# Relatório ------------------------------------------------------------------

def _texto_regra_prazo(resultado: ResultadoValidacao, prazo_padrao_dias: int | None) -> str:
    partes = []
    if C.DATA_PREVISTA in resultado.campos_presentes:
        partes.append("coluna «Data Prevista» do ficheiro")
    if prazo_padrao_dias:
        partes.append(f"prazo padrão de {prazo_padrao_dias} dias após a data da encomenda"
                      + (" (quando não há data prevista)" if partes else ""))
    return "; ".join(partes) if partes else "nenhuma — os atrasos não foram calculados"


def _percentagem(parte: int, total: int) -> float | None:
    return parte / total if total else None


def exportar_relatorio(
    resultado: ResultadoValidacao,
    tabela: pd.DataFrame,
    indicadores: Indicadores,
    nome_origem: str,
    prazo_padrao_dias: int | None = None,
) -> bytes:
    """Relatório em Excel: resumo, estados, atrasadas, problemas e lista completa.

    ``tabela`` deve ser a tabela devolvida por ``aplicar_prazos`` (todas as linhas)
    e ``indicadores`` os indicadores calculados sobre essa mesma tabela.
    """
    ind = indicadores
    livro = Workbook()

    # 1. Resumo -------------------------------------------------------------
    ws = livro.active
    ws.title = "Resumo"
    _escrever(ws, 1, 1, "LogiFlow — Relatório de encomendas").font = _FONTE_TITULO
    linha = _escrever_pares(ws, 3, [
        ("Ficheiro de origem", nome_origem),
        ("Gerado em", _agora()),
        ("Data de referência", resultado.data_referencia),
        ("Prazo usado para atrasos", _texto_regra_prazo(resultado, prazo_padrao_dias)),
        ("Âmbito", f"Todas as {ind.total} linhas importadas (incluindo {ind.com_erros} com erros)"),
    ])

    _escrever(ws, linha + 1, 1, "Indicadores").font = _FONTE_SECCAO
    if ind.atrasadas is None:
        atrasadas, nota_atrasadas = "Não disponível", (
            "Sem data prevista nem prazo padrão: não é possível determinar atrasos.")
    else:
        atrasadas = ind.atrasadas
        nota_atrasadas = (f"Há {ind.pendentes_sem_prazo} pendente(s) sem dados de prazo, que não "
                          "foram avaliadas." if ind.pendentes_sem_prazo else "")
    linhas_ind = [
        ("Total de encomendas", ind.total, None, ""),
        ("Entregues", ind.entregues, _percentagem(ind.entregues, ind.total),
         "Só conta como entregue o estado «Entregue» indicado no ficheiro."),
        ("Pendentes", ind.pendentes, _percentagem(ind.pendentes, ind.total),
         "Estados «Pendente» e «Em trânsito»."),
        ("Atrasadas", atrasadas,
         _percentagem(ind.atrasadas, ind.total) if ind.atrasadas is not None else None, nota_atrasadas),
        ("Entregues com atraso", ind.entregues_com_atraso, _percentagem(ind.entregues_com_atraso, ind.total),
         "Entregues depois do prazo (só quando há data de entrega e prazo)."),
        ("Canceladas", ind.canceladas, _percentagem(ind.canceladas, ind.total), ""),
        ("Devolvidas", ind.devolvidas, _percentagem(ind.devolvidas, ind.total), ""),
        ("Estado não reconhecido", ind.estado_nao_reconhecido,
         _percentagem(ind.estado_nao_reconhecido, ind.total), "Não contam como entregues nem pendentes."),
        ("Sem estado", ind.sem_estado, _percentagem(ind.sem_estado, ind.total), ""),
        ("Linhas com erros", ind.com_erros, _percentagem(ind.com_erros, ind.total),
         "Ver folha «Problemas»."),
    ]
    inicio = linha + 2
    fim = _escrever_tabela(ws, ["Indicador", "Valor", "% do total", "Nota"], linhas_ind,
                           linha_inicial=inicio, filtro=False)
    for k in range(inicio + 1, fim + 1):
        ws.cell(row=k, column=3).number_format = "0.0%"

    _escrever(ws, fim + 2, 1, "Qualidade dos dados").font = _FONTE_SECCAO
    problemas = resultado.problemas
    if problemas.empty:
        _escrever(ws, fim + 3, 1, "Não foram encontrados problemas.")
    else:
        contagem = (problemas.groupby(["Tipo", "Gravidade"]).size().unstack(fill_value=0)
                    .reindex(columns=[ERRO, AVISO], fill_value=0))
        _escrever_tabela(ws, ["Tipo de problema", "Erros", "Avisos"],
                         ([tipo, int(r[ERRO]), int(r[AVISO])] for tipo, r in contagem.iterrows()),
                         linha_inicial=fim + 3, filtro=False)
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 70

    # 2. Por estado -----------------------------------------------------------
    ws = livro.create_sheet("Por estado")
    _escrever_tabela(ws, ["Estado", "Encomendas", "% do total"], (
        [estado, n, _percentagem(n, ind.total)] for estado, n in ind.por_estado.items()
    ), filtro=False)
    for k in range(2, len(ind.por_estado) + 2):
        ws.cell(row=k, column=3).number_format = "0.0%"
    linha = len(ind.por_estado) + 3
    contagem_sit = tabela[COL_SITUACAO].value_counts()
    _escrever_tabela(ws, ["Situação face ao prazo", "Encomendas"], (
        [s, int(contagem_sit.get(s, 0))] for s in SITUACOES
    ), linha_inicial=linha, filtro=False)

    # 3. Atrasadas ------------------------------------------------------------
    ws = livro.create_sheet("Atrasadas")
    atrasadas_df = tabela[tabela[COL_SITUACAO] == SIT_ATRASADA].sort_values(COL_DIAS_ATRASO, ascending=False)
    if ind.atrasadas is None:
        _escrever(ws, 1, 1, "Não é possível determinar atrasos: o ficheiro não tem «Data Prevista» "
                            "e não foi definido um prazo padrão.")
    else:
        _escrever_tabela(ws, ["N.º Encomenda", "Cliente", "Destino", "Estado", "Data", "Prazo",
                              "Origem do prazo", "Dias de atraso", ROTULO_LINHA],
                         _linhas(atrasadas_df, [C.NUMERO, C.CLIENTE, C.DESTINO, C.ESTADO, C.DATA, COL_PRAZO,
                                                COL_ORIGEM_PRAZO, COL_DIAS_ATRASO, COL_LINHA]))

    # 4. Problemas ------------------------------------------------------------
    ws = livro.create_sheet("Problemas")
    _escrever_tabela(ws, list(problemas.columns), problemas.itertuples(index=False, name=None))

    # 5. Lista completa -------------------------------------------------------
    ws = livro.create_sheet("Todas as encomendas")
    campos = _campos_exportados(resultado)
    cabecalho = ([ROTULO_LINHA] + [C.rotulo(c) for c in campos]
                 + ["Prazo", "Situação face ao prazo", "Dias de atraso", "Erros", "Avisos"])
    _escrever_tabela(ws, cabecalho, _linhas(
        tabela, [COL_LINHA, *campos, COL_PRAZO, COL_SITUACAO, COL_DIAS_ATRASO, COL_ERROS, COL_AVISOS]))

    livro.properties.creator = f"LogiFlow {__version__}"
    return _gravar(livro)


# Relatório de desempenho -----------------------------------------------------

def _preparar_impressao(livro: Workbook) -> None:
    """Impressão (e PDF) em horizontal, com cada folha ajustada à largura da página."""
    for ws in livro.worksheets:
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True


def _plural(n: int, singular: str, plural: str) -> str:
    return f"{n} {singular if n == 1 else plural}"


def _texto_dias(valor: float | None) -> str:
    return f"{valor:.1f}".replace(".", ",") if valor is not None else "—"


def _tabela_df(ws: Worksheet, df: pd.DataFrame, formatos: dict[str, str], linha_inicial: int = 1,
               cabecalho: Sequence[str] | None = None) -> int:
    """Escreve um DataFrame como tabela formatada; ``formatos`` usa os nomes das colunas."""
    colunas = list(df.columns)
    ultima = _escrever_tabela(ws, list(cabecalho or colunas), _linhas(df, colunas), linha_inicial=linha_inicial)
    _formatar_colunas(ws, linha_inicial + 1, ultima,
                      {colunas.index(c) + 1: f for c, f in formatos.items() if c in colunas})
    return ultima


def exportar_desempenho(
    resultado: ResultadoValidacao,
    tabela: pd.DataFrame,
    nome_origem: str,
    prazo_padrao_dias: int | None = None,
    periodo: str = D.SEMANA,
    cliente: str | None = None,
) -> bytes:
    """Relatório de desempenho operacional (.xlsx), de todos os clientes ou de um só.

    ``tabela`` é a tabela devolvida por ``aplicar_prazos`` (todas as linhas). Com
    ``cliente``, o relatório é filtrado **antes** de qualquer cálculo: não contém
    dados de outros clientes, o que permite enviá-lo ao próprio cliente.
    """
    if cliente is not None:
        tabela = D.filtrar_cliente(tabela, cliente)
    r = D.resumir(tabela)
    formatos = {D.TAXA: FORMATO_PERCENTAGEM, D.TEMPO_MEDIO: FORMATO_DIAS}
    livro = Workbook()

    # 1. Resumo -------------------------------------------------------------
    ws = livro.active
    ws.title = "Resumo"
    _escrever(ws, 1, 1, "LogiFlow — Relatório de desempenho").font = _FONTE_TITULO
    periodo_dados = (f"{formatar_data(r.data_inicio)} a {formatar_data(r.data_fim)}"
                     if r.data_inicio else "—")
    base = _plural(r.base, "encomenda sem erros", "encomendas sem erros")
    if r.excluidas_com_erros:
        base += " (" + _plural(r.excluidas_com_erros, "linha com erros ficou de fora",
                               "linhas com erros ficaram de fora") + ")"
    linha = _escrever_pares(ws, 3, [
        ("Âmbito", f"Cliente «{cliente}»" if cliente is not None else "Todos os clientes"),
        ("Ficheiro de origem", nome_origem),
        ("Gerado em", _agora()),
        ("Data de referência", resultado.data_referencia),
        ("Datas das encomendas", periodo_dados),
        ("Prazo usado", _texto_regra_prazo(resultado, prazo_padrao_dias)),
        ("Base de cálculo", base),
    ])

    _escrever(ws, linha + 1, 1, "Indicadores").font = _FONTE_SECCAO
    if r.taxa_no_prazo is None:
        taxa, detalhe_taxa = "Não disponível", "Nenhuma entrega tem data de entrega e prazo."
    else:
        taxa = r.taxa_no_prazo
        detalhe_taxa = f"{r.no_prazo} de {r.avaliadas} entregas avaliadas"
        if r.amostra_pequena:
            detalhe_taxa += f" — poucas entregas (menos de {D.AMOSTRA_MINIMA}): use como indicação"
    if r.atrasadas is None:
        atrasadas, detalhe_atrasadas = "Não disponível", "As pendentes não têm prazo."
    else:
        atrasadas = r.atrasadas
        detalhe_atrasadas = (f"{r.pendentes_sem_prazo} pendente(s) sem prazo não avaliadas"
                             if r.pendentes_sem_prazo else "")
    detalhe_tempo = (f"Mediana: {_texto_dias(r.tempo_mediano_dias)} dias; {r.com_tempo} entregas com as duas datas"
                     if r.com_tempo else "Nenhuma entrega tem data da encomenda e data de entrega.")
    if r.incoerentes:
        detalhe_tempo += f"; {r.incoerentes} com entrega antes da encomenda ficaram de fora"
    indicadores = [
        ("Taxa de entregas no prazo", taxa, detalhe_taxa),
        ("Entregas com atraso", r.com_atraso,
         f"Atraso médio: {_texto_dias(r.atraso_medio_dias)} dias" if r.com_atraso else ""),
        ("Tempo médio de entrega (dias)", r.tempo_medio_dias if r.com_tempo else "Não disponível", detalhe_tempo),
        ("Entregas efetuadas", r.entregues,
         f"{r.entregues_sem_data} sem data de entrega" if r.entregues_sem_data else ""),
        ("Pendentes", r.pendentes, "Estados «Pendente» e «Em trânsito»."),
        ("Pendentes atrasadas", atrasadas, detalhe_atrasadas),
    ]
    inicio = linha + 2
    fim = _escrever_tabela(ws, ["Indicador", "Valor", "Detalhe"], indicadores, linha_inicial=inicio, filtro=False)
    if r.taxa_no_prazo is not None:
        ws.cell(row=inicio + 1, column=2).number_format = FORMATO_PERCENTAGEM
    if r.com_tempo:
        ws.cell(row=inicio + 3, column=2).number_format = FORMATO_DIAS

    _escrever_notas(ws, fim + 2, "Como foi calculado", [
        "Só entram as encomendas sem erros (as mesmas dos «dados limpos»): linhas duplicadas ou com "
        "datas inválidas distorceriam os resultados.",
        "Uma entrega é avaliada quando tem o estado «Entregue», data de entrega e prazo. Entregar no "
        "próprio dia do prazo conta como «no prazo».",
        "Tempo de entrega = data de entrega − data da encomenda, em dias.",
        f"Taxas com menos de {D.AMOSTRA_MINIMA} entregas avaliadas são assinaladas como pouco fiáveis.",
        "Clientes e destinos escritos de formas diferentes (maiúsculas, acentos, pontuação) são agrupados.",
    ], largura=3)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 62

    # 2. Por período -------------------------------------------------------
    ws = livro.create_sheet("Por período")
    por_periodo = D.por_periodo(tabela, periodo).drop(columns=D.INICIO_PERIODO)
    cabecalho = [periodo if c == D.PERIODO else c for c in por_periodo.columns]
    _tabela_df(ws, por_periodo, formatos, cabecalho=cabecalho)

    # 3. Por cliente (só no relatório de todos os clientes) e por destino ---
    if cliente is None:
        _tabela_df(livro.create_sheet("Por cliente"), D.por_grupo(tabela, C.CLIENTE), formatos)
    _tabela_df(livro.create_sheet("Por destino"), D.por_grupo(tabela, C.DESTINO), formatos)

    # 4. Pendentes por antiguidade -------------------------------------------
    _tabela_df(livro.create_sheet("Pendentes por antiguidade"),
               D.antiguidade_pendentes(tabela, resultado.data_referencia), {})

    # 5. Entregas com atraso (a lista que sustenta a taxa) --------------------
    ws = livro.create_sheet("Entregas com atraso")
    _escrever_tabela(ws, ["N.º Encomenda", "Cliente", "Destino", "Data", "Prazo", "Data de Entrega",
                          "Dias de atraso", ROTULO_LINHA],
                     _linhas(D.entregas_com_atraso(tabela), [C.NUMERO, C.CLIENTE, C.DESTINO, C.DATA, COL_PRAZO,
                                                            C.DATA_ENTREGA, COL_DIAS_ATRASO, COL_LINHA]))

    _preparar_impressao(livro)
    livro.properties.creator = f"LogiFlow {__version__}"
    return _gravar(livro)


# Exceções ---------------------------------------------------------------------

def exportar_excecoes(
    excecoes: pd.DataFrame,
    nome_origem: str,
    data_referencia: dt.date,
    filtros: Sequence[tuple[str, str]] = (),
) -> bytes:
    """Excel com as exceções recebidas — na aplicação, as que estão visíveis no painel
    depois de todos os filtros, pela mesma ordem.

    Não acrescenta linhas, clientes nem totais que não venham em ``excecoes``: o que não
    está visível não é exportado. ``filtros`` descreve os filtros aplicados (pares
    rótulo/valor), para ficarem registados no ficheiro.
    """
    tabela = E.visiveis(excecoes)
    totais = E.contagens(excecoes)
    livro = Workbook()

    # 1. Exceções -----------------------------------------------------------
    ws = livro.active
    ws.title = "Exceções"
    cabecalho = [ROTULO_LINHA if c == E.LINHA else c for c in tabela.columns]
    _escrever_tabela(ws, cabecalho, _linhas(tabela, list(tabela.columns)))
    if tabela.empty:
        _escrever(ws, 3, 1, "Não há exceções com os filtros aplicados.")
    for coluna in (E.MOTIVOS, E.AVISOS):
        letra = get_column_letter(list(tabela.columns).index(coluna) + 1)
        ws.column_dimensions[letra].width = 60
        for celula in ws[letra][1:]:  # texto longo: quebra de linha, também na impressão
            celula.alignment = Alignment(wrap_text=True, vertical="top")

    # 2. Informação ---------------------------------------------------------
    ws = livro.create_sheet("Informação")
    _escrever(ws, 1, 1, "LogiFlow — Exceções").font = _FONTE_TITULO
    linha = _escrever_pares(ws, 3, [
        ("Ficheiro de origem", nome_origem),
        ("Gerado em", _agora()),
        ("Data de referência", data_referencia),
        ("Encomendas exportadas", totais["encomendas"]),
        *filtros,
    ])
    _escrever(ws, linha + 1, 1, "Por prioridade").font = _FONTE_SECCAO
    fim = _escrever_tabela(ws, ["Prioridade", "Encomendas"],
                           list(totais["por_prioridade"].items()), linha_inicial=linha + 2, filtro=False)
    _escrever(ws, fim + 2, 1, "Por categoria").font = _FONTE_SECCAO
    fim = _escrever_tabela(ws, ["Categoria", "Encomendas", "Prioridade", "Regra"],
                           [(c, n, E.PRIORIDADE_DA_CATEGORIA[c], E.REGRAS[c])
                            for c, n in totais["por_categoria"].items()],
                           linha_inicial=fim + 3, filtro=False)
    _escrever_notas(ws, fim + 2, "Notas", [
        "Uma encomenda pode estar em várias categorias: os totais por categoria não se somam.",
        "Cada linha do ficheiro de origem aparece uma só vez, com todos os motivos e a prioridade mais alta.",
        "Prioridade: Alta = atrasada; Média = prazo próximo ou pendente há demasiado tempo; "
        "Baixa = só problemas de qualidade de dados.",
        "Os prazos calculados com o prazo padrão estão identificados na coluna «Origem do prazo».",
    ])
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 80

    _preparar_impressao(livro)
    livro.properties.creator = f"LogiFlow {__version__}"
    return _gravar(livro)
