"""Gera os ficheiros de exemplo em ``dados_exemplo/``.

Todos os dados são fictícios: os nomes de empresas foram inventados e não há
pessoas, moradas, contribuintes ou contactos. Os destinos são apenas nomes de
localidades portuguesas.

O gerador é determinístico (usa sempre a mesma "semente"), por isso volta a
produzir exatamente as mesmas encomendas. Execução, a partir da pasta logiflow:

    python scripts/gerar_exemplo.py
"""

from __future__ import annotations

import csv
import datetime as dt
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

PASTA = Path(__file__).resolve().parent.parent / "dados_exemplo"
DATA_REFERENCIA = dt.date(2026, 10, 20)  # usada na demonstração (ver README)

CABECALHO = ["N.º Encomenda", "Cliente", "Data", "Destino", "Estado",
             "Data Prevista", "Data de Entrega", "Observações"]

CLIENTES = [
    "Mercearia Vale Azul, Lda.", "Têxteis Ribeira Alta, S.A.", "Papelaria Sete Ventos",
    "Café Aurora Boreal", "Oficina Rota do Sul", "Livraria Quinta das Letras",
    "Florista Jardim Imaginário", "Ferragens Ponte Nova", "Vinhos Encosta Serena, Lda.",
    "Padaria Moinho Dourado", "Eletro Faísca, Lda.", "Cerâmicas Barro Fino",
]
DESTINOS = [
    "Lisboa", "Porto", "Braga", "Coimbra", "Aveiro", "Faro", "Évora", "Viseu", "Leiria",
    "Setúbal", "Guimarães", "Funchal", "Ponta Delgada", "Viana do Castelo", "Castelo Branco",
]
OBSERVACOES = ["", "", "", "", "Frágil", "Entregar de manhã", "Cais de carga 2",
               "Contactar a receção à chegada", "2 paletes"]


def gerar() -> list[list[object]]:
    aleatorio = random.Random(2026)
    linhas: list[list[object]] = []
    inicio = dt.date(2026, 9, 1)
    for i in range(60):
        data = inicio + dt.timedelta(days=int(i * 0.82))
        prevista = data + dt.timedelta(days=aleatorio.choice([2, 3, 3, 4, 5]))
        entrega: dt.date | str = ""
        if prevista <= DATA_REFERENCIA - dt.timedelta(days=6):
            sorteio = aleatorio.random()
            if sorteio < 0.06:
                estado = "Cancelada"
            elif sorteio < 0.09:
                estado = "Devolvida"
            else:
                estado = "Entregue"
                # A maioria chega dentro do prazo; algumas chegam 1 ou 2 dias depois.
                desvio = aleatorio.choice([-2, -1, -1, 0, 0, 0, 0, 0, 1, 2])
                entrega = max(data + dt.timedelta(days=1), prevista + dt.timedelta(days=desvio))
        else:
            estado = aleatorio.choice(["Pendente", "Em trânsito"])
        linhas.append([
            f"ENC-26-{i + 1:04d}", aleatorio.choice(CLIENTES), data, aleatorio.choice(DESTINOS),
            estado, prevista, entrega, aleatorio.choice(OBSERVACOES),
        ])

    # Problemas colocados de propósito, para demonstrar a validação --------
    por_entregar = [i for i, l in enumerate(linhas) if l[4] in ("Pendente", "Em trânsito")]
    entregues = [i for i, l in enumerate(linhas) if l[4] == "Entregue"]
    linhas[7][1] = ""                          # cliente vazio
    linhas[15][3] = ""                         # destino vazio
    linhas[22][0] = linhas[21][0]              # n.º repetido (dados diferentes)
    linhas[30][2] = "31/02/2026"               # data que não existe
    linhas[38][5] = "brevemente"               # data prevista ilegível
    linhas[entregues[-1]][4] = "A confirmar"   # estado não reconhecido (aviso)
    linhas[entregues[-1]][6] = ""
    i = por_entregar[1]                        # tem data de entrega mas não está entregue (aviso)
    linhas[i][6] = linhas[i][2] + dt.timedelta(days=2)
    linhas[por_entregar[2]][4] = ""            # estado vazio
    linhas[59] = list(linhas[58])              # linha repetida por inteiro
    # Variações de escrita que a aplicação uniformiza sem as tratar como erro.
    linhas[entregues[5]][4] = "entregue"
    linhas[por_entregar[4]][4] = "em transito"
    return linhas


def _texto_data(valor: object) -> object:
    return valor.strftime("%d/%m/%Y") if isinstance(valor, dt.date) else valor


def gravar_csv(linhas: list[list[object]], caminho: Path) -> None:
    # UTF-8 com BOM e ";" — é o que o Excel em português abre corretamente.
    with caminho.open("w", encoding="utf-8-sig", newline="") as f:
        escritor = csv.writer(f, delimiter=";")
        escritor.writerow(CABECALHO)
        for linha in linhas:
            escritor.writerow([_texto_data(v) for v in linha])


def gravar_xlsx(linhas: list[list[object]], caminho: Path) -> None:
    livro = Workbook()
    folha = livro.active
    folha.title = "Encomendas"
    folha.append(CABECALHO)
    for celula in folha[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="1F4E79")
    for linha in linhas:
        folha.append(linha)
    for linha in folha.iter_rows(min_row=2):
        for celula in linha:
            if isinstance(celula.value, dt.date):
                celula.number_format = "DD/MM/YYYY"
    for letra, largura in zip("ABCDEFGH", [16, 30, 12, 18, 14, 14, 16, 30]):
        folha.column_dimensions[letra].width = largura
    folha.freeze_panes = "A2"
    fixo = dt.datetime(2026, 10, 20, 9, 0)
    livro.properties.creator = "LogiFlow (dados fictícios)"
    livro.properties.created = fixo
    livro.save(caminho)


def main() -> None:
    PASTA.mkdir(exist_ok=True)
    linhas = gerar()
    gravar_csv(linhas, PASTA / "encomendas_exemplo.csv")
    gravar_xlsx(linhas, PASTA / "encomendas_exemplo.xlsx")
    print(f"Gerados {len(linhas)} registos em {PASTA}")


if __name__ == "__main__":
    main()
